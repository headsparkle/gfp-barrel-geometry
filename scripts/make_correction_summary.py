#!/usr/bin/env python3
"""Generate the cover letter to the editor accompanying the revised manuscript
(JCIM ci-2026-01606c), summarizing the response to the reviewers and the
post-submission quantum-yield re-curation. A detailed point-by-point response is
a separate document."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

def H(text, size=13, space_before=12):
    p = doc.add_paragraph(); p.space_before = Pt(space_before)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
    return p

def body(text):
    p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def bullet(text):
    p = doc.add_paragraph(text, style='List Bullet'); p.paragraph_format.space_after = Pt(4)
    return p

# ---- Title block ----
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('Cover Letter — Revised Manuscript'); r.bold = True; r.font.size = Pt(16)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Manuscript ID ci-2026-01606c\n'
                '“Barrel Shape and Chromophore Rigidity Predict Fluorescent-Protein Photophysics”\n'
                'L. P. Begg, M. L. Mason, M. Zimmer')
r.italic = True; r.font.size = Pt(11)
doc.add_paragraph()

# ---- Salutation / opening ----
body('Dear Professor He,')
body(
    'Thank you for the opportunity to submit a revised manuscript, and please convey '
    'our thanks to the three reviewers for their careful and constructive assessments. '
    'We have made major revisions that address each of the concerns raised; a detailed '
    'point-by-point response accompanies this letter. The principal changes are '
    'summarized below.')

# ---- Disclosure ----
H('Post-submission data re-curation (disclosure)')
body(
    'Independently of the reviews, our continued data-quality review identified an '
    'artifact in the quantum-yield (QY) annotations of the submitted version: public '
    'FP databases propagate the wild-type avGFP quantum yield (0.79) to many entries by '
    'keyword inheritance, so 82% of annotated structures carried 0.79 and only 14 '
    'distinct values were present. We re-curated QY directly from FPbase (by PDB '
    'identifier, else chain-A sequence at ≥ 99% identity), recovering 354 structures '
    'spanning 0.0001–0.97. This is fully scripted and reproducible, and it happens to '
    'address several reviewer concerns about the QY data (Reviewers 2 and 3). The '
    'barrel-shape/emission results and the AlphaFold comparison are unchanged; the QY '
    'analysis is refined and, we believe, strengthened.')

# ---- Reviewer responses ----
H('1. Sampling imbalance and pseudoreplication (Reviewers 1, 2, 3)')
body(
    'We agree this was the central weakness and have addressed it directly with a new '
    'robustness analysis (Table S7).')
bullet(
    'Granular per-protein identity: each structure is now assigned its matched FPbase '
    'entry (or chain-A sequence cluster), giving 430 unique proteins in the canonical '
    'cohort rather than the coarse count used previously — a much more stringent '
    'pseudoreplication control.')
bullet(
    'Under collapse to one entry per protein the key correlations are preserved or '
    'strengthened (redundancy attenuated, not inflated them); they also hold in a '
    'monomer-only re-run (e.g., emission vs minor axis ρ = –0.46), ruling out '
    'oligomeric-packing artifacts, and the emission–geometry relationship holds within '
    'the green class alone (ρ = –0.20, p = 3 × 10⁻⁵).')
bullet(
    'The blue (n = 10) and orange (n = 11) classes are too small for within-class '
    'inference; we now report them descriptively only and explicitly restrict '
    'within-class quantitative claims to the well-sampled green, red, cyan, and yellow '
    'classes (revised Limitations).')

H('2. Correlation versus causation (Reviewer 3)')
body(
    'We appreciate the importance of causal evidence. As a computational and structural '
    'bioinformatics study — the scope of the work and of this journal — new '
    'site-directed-mutagenesis experiments are outside what we can appropriately '
    'undertake here, and we have therefore reframed the findings as rigorous, '
    'hypothesis-generating structural correlations, softened causal language, and added '
    'an explicit limitations statement. We also marshal the existing experimental '
    'record that bears on the hypotheses — the T203Y/T203H mutations in the GFP/YFP '
    'lineage and the I146F substitution in mTurquoise2 — and cite recent excited-state '
    'studies (Pieri et al., J. Am. Chem. Soc. 2024; Chen et al., Proc. Natl. Acad. Sci. '
    'U. S. A. 2025) that experimentally and computationally support the proposed links '
    'between ground-state geometry, rigidity, and brightness.')

H('3. Quantum-yield data, metric, and model performance (Reviewers 2, 3)')
bullet(
    'QY data availability is much improved by the re-curation (354 structures vs 321), '
    'and all QY statistics are now reported per unique protein.')
bullet(
    'We adopted a symmetry-folded distance-from-planar metric for ground-state '
    'geometry, correcting a case in which an unfolded twist sum mis-scored ring-flipped '
    'chromophores (e.g., near-planar mScarlet).')
bullet(
    'The modest model R² (~0.18) is now explained as expected for a structure-only '
    'predictor: quantum yield is set at the electronic-structure level (excited-state '
    'surface topology, conical-intersection accessibility, protonation, charge transfer, '
    'dark-state pathways) not encoded in static crystal geometry. Our models identify '
    'which structural features carry reproducible signal rather than serving as '
    'quantitative predictors.')

H('4. Terminology, definitions, and presentation (Reviewers 1, 2)')
body(
    'We now define “fluorescence quantum yield” on first use, spell out principal '
    'component analysis (PCA), clarify that the dashed lines in Figure 1 are '
    'guides-to-the-eye least-squares fits while the reported Spearman ρ quantifies each '
    'association, and add the requested citations. Remaining presentation items raised '
    'by the reviewers (figure call-outs and ordering, additional figures, and '
    'typographic conventions such as italic cis/trans) are addressed in the revised '
    'manuscript and detailed in the point-by-point response.')

# ---- before/after table ----
H('Summary of the principal numerical changes')
tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Light Grid Accent 1'
for c, txt in zip(tbl.rows[0].cells, ['Quantity', 'As submitted', 'Revised']):
    c.paragraphs[0].add_run(txt).bold = True
rows = [
    ('Quantum-yield values in dataset', '14 distinct (82% = avGFP 0.79)',
     '354 curated from FPbase (0.0001–0.97)'),
    ('Unique proteins (pseudoreplication control)', 'coarse name grouping',
     '430 (matched FPbase entry / chain-A sequence)'),
    ('Primary QY conclusion', 'B-factor ratio the single strongest predictor (ρ = –0.443)',
     'Two independent signals — chromophore-to-barrel B-factor ratio (ρ = –0.49) and '
     'ground-state planarity (ρ = –0.42); barrel size not predictive'),
    ('QY aggregation', 'per crystal structure', 'per unique protein'),
    ('Random-Forest OOB R² (QY)', '0.42', '0.18 (explained)'),
    ('Barrel shape vs emission wavelength', 'ρ = –0.329 (minor axis)', 'unchanged'),
]
for a, b, c in rows:
    cells = tbl.add_row().cells
    cells[0].paragraphs[0].add_run(a)
    cells[1].paragraphs[0].add_run(b)
    cells[2].paragraphs[0].add_run(c)
doc.add_paragraph()

# ---- closing ----
body(
    'We believe these revisions substantially strengthen the manuscript while keeping '
    'the beta-barrel geometry the focus of the study. We would be glad to provide any '
    'further clarification the reviewers or editor may require.')
body('Sincerely,')
body('Marc Zimmer, on behalf of the authors')

out = os.path.join(REPO, 'Manuscript_Correction_Summary.docx')
doc.save(out)
print('Saved:', out)
